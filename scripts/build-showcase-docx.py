"""Headless twin of the add-in's test paper: build the RevTeX-style layout with
python-docx from the outline JSON on stdin (see addin/src/testpaper.js paperOutline),
leaving equations as placeholders for `mjx-docx process`.
Usage: node scripts/make-showcase-paper.mjs  (this script is called by it)."""
import json, sys
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches

ALIGN = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'centered': WD_ALIGN_PARAGRAPH.CENTER,
         'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justified': WD_ALIGN_PARAGRAPH.JUSTIFY}

def set_cols(section, n):
    sectPr = section._sectPr
    cols = sectPr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols'); sectPr.append(cols)
    cols.set(qn('w:num'), str(n)); cols.set(qn('w:space'), '288')

def apply_fmt(p, fmt):
    pf = p.paragraph_format
    if not fmt: return
    if fmt.get('align'): p.alignment = ALIGN[fmt['align']]
    for key, attr in [('firstLineIndent', 'first_line_indent'), ('leftIndent', 'left_indent'),
                      ('rightIndent', 'right_indent'), ('spaceBefore', 'space_before'), ('spaceAfter', 'space_after')]:
        if fmt.get(key) is not None: setattr(pf, attr, Pt(fmt[key]))
    if fmt.get('lineSpacing') is not None: pf.line_spacing = Pt(fmt['lineSpacing'])

def style_run(r, fmt):
    if not fmt: return
    if fmt.get('font'): r.font.name = fmt['font']
    if fmt.get('size'): r.font.size = Pt(fmt['size'])
    r.font.bold = bool(fmt.get('bold')); r.font.italic = bool(fmt.get('italic'))

outline = json.load(sys.stdin)
out = sys.argv[1]
doc = Document()
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'; normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(0)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'): setattr(sec, side, Inches(0.75))

for i, b in enumerate(outline):
    if b['kind'] == 'section':
        # The add-in's sectPr paragraph applies `cols` to the section that ENDS
        # there. python-docx: a new continuous section starts after this point,
        # so set the columns on the section that just closed.
        if i == len(outline) - 1:
            set_cols(doc.sections[-1], b['cols'])
        else:
            doc.add_section(WD_SECTION.CONTINUOUS)
            set_cols(doc.sections[-2], b['cols'])
        continue
    p = doc.add_paragraph()
    apply_fmt(p, b.get('fmt'))
    r = p.add_run(b['text'])
    style_run(r, b.get('fmt'))
doc.save(out)
print(out)
