"""Generate deploy/Recorde-TestPage.docx — a one-page manual test document.

A short physics-style text with numbered insertion points. Each point names the
exact pane settings to use, so working through the page exercises every option
of the add-in: modes, numbering styles, fonts, size modes, colour, macros,
palette, examples, click-to-edit, renumber, the equation list and the theme.

Run:  micromamba run -n sci-base python scripts/make-test-page.py
"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(11)

def h(text, level=1):
    doc.add_heading(text, level=level)

def p(text):
    return doc.add_paragraph(text)

def slot(n, settings, latex, note=''):
    """A visible insertion point: put the cursor on the ⟨…⟩ marker, delete it, insert."""
    para = doc.add_paragraph()
    r = para.add_run(f'⟨{n}⟩ ')
    r.bold = True; r.font.color.rgb = RGBColor(0xB5, 0x32, 0x1E)
    r = para.add_run(settings + ' — ')
    r.italic = True
    r = para.add_run(latex)
    r.font.name = 'Consolas'; r.font.size = Pt(9.5)
    if note:
        r = para.add_run('  ' + note); r.font.size = Pt(9); r.font.color.rgb = RGBColor(0x61, 0x61, 0x61)
    return para

h('Recorde test page', 0)
p('Work through the markers ⟨n⟩ in order. For each: click on the marker line, delete the marker text, '
  'set the pane as described, paste the LaTeX, Insert. Tick the result in the table at the end. '
  'Pane settings not mentioned stay at their defaults (Termes, Match selection, black).')

h('1. Inline equations in running text', 1)
p('Consider a particle of mass m in a potential V(x). Its energy is conserved, ')
slot(1, 'Inline · Termes · Match selection', r'E = \frac{p^2}{2m} + V(x)',
     'baseline should sit on the text line')
p('and for a harmonic well the frequency follows from the curvature at the minimum, ')
slot(2, 'Inline · New Computer Modern · Match body', r'\omega = \sqrt{V''(x_0)/m}',
     'different font than ⟨1⟩')
p('The imaginary unit and Euler’s number come from the document preamble: ')
slot(3, 'Inline · Termes · Fixed 14 pt · colour dark red (#8B0000)', r'\E^{\I\pi} + 1 = 0',
     'larger + coloured; uses \\E and \\I macros')

h('2. Display equations, unnumbered', 1)
p('The propagator of the free particle in one dimension is')
slot(4, 'Display · STIX Two · Match body', r'K(x_f,t_f;x_i,t_i) = \sqrt{\frac{m}{2\pi\I\hbar\,(t_f-t_i)}}\;\exp\!\left[\frac{\I m (x_f-x_i)^2}{2\hbar (t_f-t_i)}\right]',
     'centred on its own paragraph; the prose stays left-aligned')

h('3. Numbering: inline style', 1)
p('Set Numbering to “Inline (after equation)” and tick “Number this equation”.')
slot(5, 'Display · numbered (inline style)', r'\int_{-\infty}^{\infty} e^{-x^2}\,dx = \sqrt{\pi}', 'expect (1)')
slot(6, 'Display · numbered (inline style) · Pagella', r'x = \frac{-b \pm \sqrt{b^2 - 4ac}}{2a}', 'expect (2)')

h('4. Numbering: flush-right table style', 1)
p('Switch Numbering to “Flush right (table)”.')
slot(7, 'Display · numbered (table style)', r'\VNabla \cdot \vb{E} = \frac{\rho}{\varepsilon_0}',
     'expect a borderless row: equation centred, (3) at the right margin')
slot(8, 'Display · numbered (table style) · Asana Math', r'\HOp{H}\,\ket{\psi} = E\,\ket{\psi}', 'expect (4)')

h('5. Numbering: Word field style (desktop only)', 1)
p('Switch Numbering to “Word field (desktop)”.')
slot(9, 'Display · numbered (field style)', r'\Avg*{\HOp{H}} = \sum_n \Abs{c_n}^2 E_n', 'expect (5); Alt+F9 shows { SEQ equation \\* ARABIC }')
slot(10, 'Display · numbered (field style)', r'\OpComm{\HOp{x}, \HOp{p}} = \I\hbar', 'expect (6)')

h('6. Editing and renumbering', 1)
p('a) Click equation ⟨5⟩ in the document: the pane loads it, status says “Editing”, the number shows (1) in the '
  'preview. Change the integrand to e^{-x^2/2} and press Ctrl+Enter: it updates in place and keeps (1).')
p('b) Select the whole ⟨6⟩ paragraph and delete it. Click “Renumber all”: ⟨7⟩…⟨10⟩ become (2)…(5); '
  'the field equations renumber too (Ctrl+A, F9 must agree).')
p('c) Open “Equations in this document”: five entries in document order with their numbers; click the '
  'last one → Word jumps to ⟨10⟩ and the pane loads it.')
p('d) With ⟨10⟩ loaded, press Esc: the pane returns to “New equation” without touching the document.')

h('7. Tools', 1)
p('a) Examples → “Navier–Stokes”: Display mode and numbering are set; insert it here:')
slot(11, 'Examples → Navier–Stokes (as loaded)', '(from the picker)', 'expect (6) in the current style')
p('b) Symbols palette: with an empty editor click ∑, then type the limits; insert inline:')
slot(12, 'Inline · palette snippet', r'\sum_{k=1}^{N} k = \frac{N(N+1)}{2}')
p('c) Macros: add  \\newcommand{\\vecr}{\\vb{r}}  to the Macros box, then insert:')
slot(13, 'Display · uses the new macro', r'\HOp{\vb{E}}(\vecr, t)', 'must render bold r')
p('d) Type  \\frac{a}{  in the editor: the red message “Missing close brace” appears under the source; '
  'complete it and the message disappears.')
p('e) Theme: File → Account → Office Theme → Dark Gray. The pane turns dark, the preview stays white, '
  'dropdowns and checkbox match the pane. Back to Colorful.')

h('8. Results', 1)
t = doc.add_table(rows=1, cols=3)
t.style = 'Light Grid'
hdr = t.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text = 'Marker', 'OK?', 'Notes (error text, screenshot name)'
for n in range(1, 14):
    row = t.add_row().cells
    row[0].text = f'⟨{n}⟩'
for label in ('6a edit/update', '6b renumber', '6c equation list', '6d Esc', '7d error message', '7e theme'):
    row = t.add_row().cells
    row[0].text = label

out = 'deploy/Recorde-TestPage.docx'
doc.save(out)
print('wrote', out)
